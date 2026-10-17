"""
F4P direct fetchers (Batch 2).

The Hub used to ask the AI to find every figure for the seven non-USD currencies by web search. Some
official pages the AI cannot read well (data portals, JSON files, CSV downloads), so it quoted an older
page instead - an old IMF edition, a stale daily yield, a past SNB bulletin. This module downloads those
figures straight from the official source, works the comparison numbers out in code, and hands them to the
weekly update as verified facts, the same way the FRED AUTO tab already does for the US dollar.

How it fits together
  fred_cot_fetcher.py (daily)   -> calls fetch_all(), writes the DIRECT AUTO tab
  f4p_weekly_update.py (Sunday) -> reads the tab with facts_from_rows(), skips anything stale, and puts the
                                   rest in each currency's prompt as figures that are already verified

Safety rules
  * Every download is its own small unit. If one fails, the others still run and the failed one is written
    to the tab as FETCH FAILED; the weekly update then searches for that figure exactly as it did before.
  * A figure that fails a sanity check (wrong size, too old, no earlier value to compare with) is treated
    as a failure, never written. A wrong number would be worse than a missing one.
  * Nothing here writes to the Hub tabs. It only builds facts.

The module has no third-party imports at the top (requests is imported only when a download is made), so the
weekly update and the tests can import it anywhere.
"""

import csv
import datetime
import io
import json
import time
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

DIRECT_AUTO_TAB_NAME = "DIRECT AUTO"
DIRECT_AUTO_HEADER = ["CURRENCY", "INDICATOR", "FACT NAME", "VALUE", "AS OF", "SOURCE", "STATUS",
                      "RETRIEVED (UTC)", "NOTE"]
RETRIEVED_FORMAT = "%Y-%m-%d %H:%M UTC"
STATUS_OK = "OK"
STATUS_FAILED = "FETCH FAILED - the Hub will search for this one itself"

# The weekly update ignores a row whose RETRIEVED time is older than this. The fetcher runs every day, so a
# row older than three days means the fetcher has been failing and the figure may be out of date.
MAX_FACT_AGE_DAYS = 3

USER_AGENT = ("Mozilla/5.0 (compatible; F4P-Hub-DataFetcher/1.0; "
              "+https://github.com/gwintlawrence/no-README)")
TIMEOUT_SECONDS = 20
TOTAL_BUDGET_SECONDS = 300      # all downloads together; whatever is left is skipped (the Hub searches instead)

YIELD_NAME = "10-Year Government Bond Yield"
DEBT_NAME = "Government Debt / GDP"
FISCAL_NAME = "Fiscal Balance"
M2_NAME = "M2 Money Supply"
BALANCE_SHEET_NAME = "Central Bank Balance Sheet / GDP"

# Freshness limits, from the FRESHNESS RULE in f4p_spec.py (a daily figure older than 7 days is not current).
YIELD_MAX_AGE_DAYS = 7
YIELD_EARLIER_GAP_DAYS = 5      # the "nearest earlier trading day" may be at most this far before the target


class FetchError(Exception):
    """A source could not be read the way this module expects. The message says why, in plain words."""


# --------------------------------------------------------------------------------------------------
# Downloading
# --------------------------------------------------------------------------------------------------

def http_get_response(url, params=None, headers=None):
    import requests   # imported here so tests and the weekly update do not need it
    head = {"User-Agent": USER_AGENT}
    head.update(headers or {})
    resp = requests.get(url, params=params, headers=head, timeout=TIMEOUT_SECONDS)
    resp.raise_for_status()
    return resp


def http_get_text(url, params=None, headers=None):
    return http_get_response(url, params=params, headers=headers).text


def http_get_bytes(url, params=None, headers=None):
    return http_get_response(url, params=params, headers=headers).content


def snippet(text, n=200):
    """Start of a response, for error messages, so a changed file layout can be diagnosed from the log."""
    text = text if isinstance(text, str) else repr(text)
    return " ".join(text[:n].split())


# --------------------------------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------------------------------

def to_decimal(value):
    """Exact decimal from a number or a string such as '3.35' or '1,234.5'. Raises FetchError otherwise."""
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value).strip().replace(",", ""))
    except (InvalidOperation, ValueError):
        raise FetchError("could not read %r as a number" % (value,))


def plain(d):
    """Decimal as plain text without an exponent (4.10 stays 4.10)."""
    return format(d, "f")


def signed(d):
    return ("+" if d > 0 else "") + plain(d)


def parse_date(text, formats):
    text = str(text).strip()
    for fmt in formats:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise FetchError("could not read %r as a date" % (text,))


def make_fact(currency, indicator, name, value, as_of, source):
    return {"currency": currency, "indicator": indicator, "name": name,
            "value": str(value), "as_of": str(as_of), "source": source}


def clean_points(points):
    """Sorted list of (date, Decimal), one per date (the last one wins), blanks dropped."""
    by_date = {}
    for d, v in points:
        if v is None or str(v).strip() in ("", ".", "-", "..", "NA", "n.a.", "N/A"):
            continue
        by_date[d] = to_decimal(v)
    return sorted(by_date.items())


def check_range(label, value, low, high):
    if not (low <= value <= high):
        raise FetchError("%s came out as %s, outside the believable range %s to %s, so it is not used"
                         % (label, plain(value), low, high))


# --------------------------------------------------------------------------------------------------
# Turning downloaded numbers into facts
# --------------------------------------------------------------------------------------------------

def yield_facts(currency, points, source, today, indicator=YIELD_NAME):
    """Latest daily 10-year yield, the value one week earlier (the nearest earlier trading day when that
    exact day has no figure), and the change between them - the figure the score guide scores."""
    pts = clean_points(points)
    if not pts:
        raise FetchError("the download contained no yield values")
    latest_d, latest_v = pts[-1]
    age = (today - latest_d).days
    if age > YIELD_MAX_AGE_DAYS:
        raise FetchError("the newest yield is dated %s, %d days before today, which is too old to use"
                         % (latest_d, age))
    check_range("the yield", latest_v, Decimal("-2"), Decimal("20"))
    target = latest_d - datetime.timedelta(days=7)
    earlier = [p for p in pts if p[0] <= target]
    if not earlier:
        raise FetchError("the download has no yield on or before %s (one week before %s)" % (target, latest_d))
    earlier_d, earlier_v = earlier[-1]
    if (target - earlier_d).days > YIELD_EARLIER_GAP_DAYS:
        raise FetchError("the nearest earlier yield is dated %s, too far before %s" % (earlier_d, target))
    change = latest_v - earlier_v
    if abs(change) > Decimal("2"):
        raise FetchError("a seven-day move of %s points looks wrong, so it is not used" % plain(change))
    return [
        make_fact(currency, indicator, indicator, plain(latest_v) + "%", latest_d, source),
        make_fact(currency, indicator,
                  indicator + " - value one week earlier (use this as the prior_value)",
                  plain(earlier_v) + "%", earlier_d, source),
        make_fact(currency, indicator,
                  indicator + " - change over the last 7 days in percentage points, latest minus one week "
                              "earlier (the score guide scores this figure)",
                  signed(change) + " points", latest_d, source),
    ]


def imf_facts(currency, indicator, values_by_year, source, today):
    """IMF World Economic Outlook figure for the last completed calendar year and the year before it,
    plus the change. The current year and later are forecasts and are never used."""
    year = today.year - 1
    prev = year - 1
    missing = [y for y in (year, prev) if values_by_year.get(y) is None]
    if missing:
        raise FetchError("the IMF has no value for %s for %s" % (" or ".join(str(y) for y in missing), currency))
    latest, earlier = values_by_year[year], values_by_year[prev]
    if indicator == DEBT_NAME:
        check_range("debt / GDP", latest, Decimal("0"), Decimal("500"))
        check_range("debt / GDP", earlier, Decimal("0"), Decimal("500"))
        shown = plain(latest) + "% of GDP (gross general government debt)"
        shown_prev = plain(earlier) + "% of GDP"
        change_name = ("%s - change in percentage points of GDP, %d minus %d (the score guide scores this figure)"
                       % (indicator, year, prev))
    else:
        check_range("fiscal balance", latest, Decimal("-60"), Decimal("60"))
        check_range("fiscal balance", earlier, Decimal("-60"), Decimal("60"))
        shown = plain(latest) + "% of GDP (net lending +, net borrowing -; a negative number is a deficit)"
        shown_prev = plain(earlier) + "% of GDP"
        change_name = ("%s - change in the budget balance in percentage points of GDP, %d minus %d "
                       "(positive = surplus grew or deficit narrowed; the score guide scores this figure)"
                       % (indicator, year, prev))
    change = latest - earlier
    return [
        make_fact(currency, indicator, indicator, shown, str(year), source),
        make_fact(currency, indicator,
                  "%s - value for the year before, %d (use this as the prior_value)" % (indicator, prev),
                  shown_prev, str(prev), source),
        make_fact(currency, indicator, change_name, signed(change) + " points", str(year), source),
    ]


def yoy_facts(currency, indicator, monthly, source, today, unit, max_age_days, what, display_divisor=None):
    """Latest monthly level, the level for the same month a year earlier, and the year-on-year change
    worked out as (latest - earlier) / earlier x 100 to one decimal (the CALCULATION RULE in the prompts).
    `monthly` is a list of ((year, month), value). When display_divisor is given (1000 turns millions into
    billions) the two levels are shown divided by it to one decimal; the per cent change is always worked out
    from the unrounded figures."""
    by_month = {}
    for key, v in monthly:
        if v is None or str(v).strip() in ("", "-", "..", "NA", "n.a."):
            continue
        by_month[key] = to_decimal(v)
    if not by_month:
        raise FetchError("the download contained no monthly values")
    latest_key = max(by_month)
    year, month = latest_key
    # last day of that month, to judge how old the figure is
    nxt = datetime.date(year + (month == 12), month % 12 + 1, 1)
    period_end = nxt - datetime.timedelta(days=1)
    age = (today - period_end).days
    if age > max_age_days:
        raise FetchError("the newest month is %04d-%02d, %d days ago, which is too old to use" % (year, month, age))
    earlier_key = (year - 1, month)
    if earlier_key not in by_month:
        raise FetchError("the download has no value for %04d-%02d (the same month a year earlier)" % earlier_key)
    latest, earlier = by_month[latest_key], by_month[earlier_key]
    if earlier == 0:
        raise FetchError("the year-earlier value is zero")
    pct = ((latest - earlier) / earlier * 100).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
    check_range("the year-on-year change", pct, Decimal("-60"), Decimal("200"))
    label = "%04d-%02d" % latest_key
    label_prev = "%04d-%02d" % earlier_key
    if display_divisor:
        shown, shown_prev = [plain((x / display_divisor).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))
                             for x in (latest, earlier)]
    else:
        shown, shown_prev = plain(latest), plain(earlier)
    return [
        make_fact(currency, indicator, "%s - %s" % (indicator, what), "%s %s" % (shown, unit), label, source),
        make_fact(currency, indicator,
                  "%s - %s for the same month a year earlier (use this as the prior_value)" % (indicator, what),
                  "%s %s" % (shown_prev, unit), label_prev, source),
        make_fact(currency, indicator,
                  "%s - year-on-year change in per cent, worked out as (latest - year earlier) / year earlier "
                  "x 100 (the score guide scores this figure)" % indicator,
                  plain(pct) + "%", label, source),
    ]


# --------------------------------------------------------------------------------------------------
# Source readers (each reads text/bytes it is given and returns points; no network inside the parsers)
# --------------------------------------------------------------------------------------------------

def parse_boc_valet(text, series_id):
    """Bank of Canada Valet API JSON -> [(date, value)]."""
    try:
        data = json.loads(text)
        observations = data["observations"]
    except (ValueError, KeyError, TypeError):
        raise FetchError("not the Valet JSON layout. Start of response: %s" % snippet(text))
    points = []
    for obs in observations:
        cell = obs.get(series_id)
        if not isinstance(cell, dict) or "v" not in cell:
            continue
        points.append((parse_date(obs.get("d", ""), ["%Y-%m-%d"]), cell["v"]))
    if not points:
        raise FetchError("the Valet response had no values for %s. Start of response: %s" % (series_id, snippet(text)))
    return points


def _looks_like_years(d):
    return isinstance(d, dict) and d and all(str(k).isdigit() and len(str(k)) == 4 for k in d)


def _find_country_years(node, country):
    if isinstance(node, dict):
        inner = node.get(country)
        if _looks_like_years(inner):
            return inner
        for child in node.values():
            found = _find_country_years(child, country)
            if found:
                return found
    elif isinstance(node, list):
        for child in node:
            found = _find_country_years(child, country)
            if found:
                return found
    return None


def parse_imf_datamapper(text, country):
    """IMF DataMapper JSON -> {year: Decimal}. Numbers are read as exact decimals; null stays missing."""
    try:
        data = json.loads(text, parse_float=Decimal)
    except ValueError:
        raise FetchError("not JSON. Start of response: %s" % snippet(text))
    years = _find_country_years(data, country)
    if not years:
        raise FetchError("no year-by-year values for %s found in the IMF response. Start of response: %s"
                         % (country, snippet(text)))
    out = {}
    for k, v in years.items():
        if v is None or isinstance(v, bool):
            continue
        try:
            out[int(k)] = to_decimal(v)
        except FetchError:
            continue
    return out


def _rows_from_csv(text, delimiter=None):
    text = text.lstrip("\ufeff")
    if delimiter is None:
        first = next((ln for ln in text.splitlines() if ln.strip()), "")
        delimiter = ";" if first.count(";") > first.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _no_html(text):
    if text.lstrip()[:15].lower().startswith(("<!doctype", "<html", "<?xml")) or "<html" in text[:300].lower():
        raise FetchError("got a web page instead of data. Start of response: %s" % snippet(text))


def parse_sdmx_csv(text):
    """Bundesbank / SDMX-CSV -> [(date, value)] using the TIME_PERIOD and OBS_VALUE columns."""
    _no_html(text)
    rows = _rows_from_csv(text)
    if not rows:
        raise FetchError("empty response")
    header = [c.strip().upper() for c in rows[0]]
    if "TIME_PERIOD" in header and "OBS_VALUE" in header:
        t_i, v_i = header.index("TIME_PERIOD"), header.index("OBS_VALUE")
    elif len(header) >= 2 and header[0] == "" and "." in header[1]:
        # The layout the Bundesbank returned in the first live run (16 Oct 2026): one header row with the series
        # key in the second column, some description rows, then data rows of date, value, flag.
        t_i, v_i = 0, 1
    else:
        raise FetchError("neither the SDMX nor the time-series layout. Start of response: %s" % snippet(text))
    points = []
    for r in rows[1:]:
        if len(r) > max(t_i, v_i):
            try:
                points.append((parse_date(r[t_i], ["%Y-%m-%d"]), r[v_i]))
            except FetchError:
                continue
    if not points:
        raise FetchError("no daily values found. Start of response: %s" % snippet(text))
    return points


BOE_DATE_FORMATS = ["%d %b %Y", "%d/%b/%Y", "%d-%b-%Y", "%d %B %Y"]


def parse_boe_iadb(text, series_code):
    """Bank of England database CSV (rows like '15 Oct 2026,4.55' under a 'DATE,<code>' header)."""
    _no_html(text)
    rows = _rows_from_csv(text, ",")
    col = None
    points = []
    for r in rows:
        if not r:
            continue
        cells = [c.strip() for c in r]
        if cells[0].upper() == "DATE":
            col = cells.index(series_code) if series_code in cells else 1
            continue
        try:
            dt = parse_date(cells[0], BOE_DATE_FORMATS)
        except FetchError:
            continue
        if len(cells) > (col or 1):
            points.append((dt, cells[col or 1]))
    if not points:
        raise FetchError("no dated values for %s found. Start of response: %s" % (series_code, snippet(text)))
    return points


def parse_mof_jgb(text):
    """Japan Ministry of Finance JGB yield table (jgbcme.csv) -> [(date, 10-year yield)]."""
    _no_html(text)
    rows = _rows_from_csv(text, ",")
    col = None
    points = []
    for r in rows:
        cells = [c.strip() for c in r]
        if not cells:
            continue
        if cells[0].lower() in ("date", "基準日") and col is None:
            for i, name in enumerate(cells):
                if name.upper().replace(" ", "") in ("10Y", "10", "10YEAR", "10YEARS", "10年"):
                    col = i
            if col is None:
                raise FetchError("no 10-year column in the header row: %s" % cells)
            continue
        if col is None or len(cells) <= col:
            continue
        try:
            dt = parse_date(cells[0], ["%Y/%m/%d", "%Y-%m-%d"])
        except FetchError:
            continue
        points.append((dt, cells[col]))
    if col is None:
        raise FetchError("could not find the header row with Date. Start of response: %s" % snippet(text))
    if not points:
        raise FetchError("no dated rows found. Start of response: %s" % snippet(text))
    return points


def parse_rba_table(text, series_ids):
    """RBA statistical table CSV: a 'Series ID' row names the columns, data rows start with a date."""
    _no_html(text)
    rows = _rows_from_csv(text, ",")
    col = None
    points = []
    for r in rows:
        cells = [c.strip() for c in r]
        if not cells:
            continue
        if cells[0].lower() == "series id":
            for i, name in enumerate(cells):
                if name in series_ids:
                    col = i
            if col is None:
                raise FetchError("series %s not in the Series ID row: %s" % (series_ids, cells[:12]))
            continue
        if col is None or len(cells) <= col:
            continue
        try:
            dt = parse_date(cells[0], ["%d-%b-%Y", "%d/%m/%Y", "%Y-%m-%d"])
        except FetchError:
            continue
        points.append((dt, cells[col]))
    if col is None:
        raise FetchError("no Series ID row found. Start of response: %s" % snippet(text))
    if not points:
        raise FetchError("no dated rows found for %s" % (series_ids,))
    return points


def parse_snb_cube(text):
    """SNB data portal cube CSV (semicolon separated, a few metadata lines, then a header starting with Date).
    Returns (dimension_names, [{'Date': ..., 'Value': ..., '<dim>': ...}])."""
    _no_html(text)
    rows = _rows_from_csv(text, ";")
    header_at = next((i for i, r in enumerate(rows) if r and r[0].strip().lower() == "date"), None)
    if header_at is None:
        raise FetchError("no header row starting with Date. Start of response: %s" % snippet(text))
    header = [c.strip() for c in rows[header_at]]
    if "Value" not in header:
        raise FetchError("no Value column in header %s" % header)
    dims = [h for h in header if h not in ("Date", "Value")]
    out = []
    for r in rows[header_at + 1:]:
        if len(r) == len(header):
            out.append(dict(zip(header, [c.strip() for c in r])))
    if not out:
        raise FetchError("the cube returned no data rows. Start of response: %s" % snippet(text))
    return dims, out


def snb_series(dims, rows, wanted, what):
    """Rows of one series: those where some dimension holds one of the `wanted` codes. If more than one
    series is left (another dimension still varies) the layout is not what was expected, so the error lists
    what each dimension holds, which is enough to correct the reader from the run log."""
    wanted = set(wanted)
    picked = [r for r in rows if any(r.get(dim) in wanted for dim in dims)]
    if not picked:
        seen = {dim: sorted({r.get(dim) for r in rows})[:30] for dim in dims}
        raise FetchError("no row with %s for %s; the cube holds %s" % (sorted(wanted), what, seen))
    leftover = [dim for dim in dims if len({r.get(dim) for r in picked}) > 1]
    # a dimension that varies only because it is the one holding the wanted code is fine
    leftover = [dim for dim in leftover if not ({r.get(dim) for r in picked} <= wanted)]
    if leftover:
        seen = {dim: sorted({r.get(dim) for r in picked})[:30] for dim in leftover}
        raise FetchError("more than one series matches %s for %s; still varying: %s" % (sorted(wanted), what, seen))
    return picked


def parse_rbnz_b2(content):
    """RBNZ B2 daily workbook -> [(date, 10-year government bond yield)] (best effort: the column is found by
    its header text, and the error lists every column heading seen if it is not found)."""
    try:
        import openpyxl
    except ImportError:
        raise FetchError("openpyxl is not installed")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:   # not an xlsx (for example an access-denied page)
        raise FetchError("not an Excel file (%s). Start of response: %s" % (exc, snippet(content[:200].decode("latin-1"))))
    ws = wb[wb.sheetnames[0]] if "Data" not in wb.sheetnames else wb["Data"]
    rows = list(ws.iter_rows(values_only=True))
    head_rows = rows[:12]
    width = max((len(r) for r in head_rows), default=0)
    labels = []
    for c in range(width):
        labels.append(" ".join(str(r[c]) for r in head_rows if c < len(r) and r[c] is not None).lower())
    cands = [i for i, lab in enumerate(labels)
             if "10" in lab and "year" in lab and any(k in lab for k in ("govern", "govt", "bond", "stock"))
             and not any(k in lab for k in ("swap", "bill", "index", "inflation"))]
    if len(cands) != 1:
        raise FetchError("found %d columns that look like the 10-year government bond yield; headings seen: %s"
                         % (len(cands), [l[:60] for l in labels[:40]]))
    col = cands[0]
    points = []
    for r in rows:
        if len(r) <= col or r[0] is None or r[col] is None:
            continue
        first = r[0]
        if isinstance(first, datetime.datetime):
            dt = first.date()
        elif isinstance(first, datetime.date):
            dt = first
        else:
            try:
                dt = parse_date(first, ["%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y"])
            except FetchError:
                continue
        points.append((dt, r[col]))
    if not points:
        raise FetchError("no dated values in the 10-year column")
    return points


# --------------------------------------------------------------------------------------------------
# Units: one entry per download. Each has the currency and indicator it feeds, so a failure can be shown.
# --------------------------------------------------------------------------------------------------

IMF_COUNTRY = {"GBP": "GBR", "JPY": "JPN", "CHF": "CHE", "CAD": "CAN", "AUD": "AUS", "NZD": "NZL"}
IMF_SERIES = {DEBT_NAME: "GGXWDG_NGDP", FISCAL_NAME: "GGXCNL_NGDP"}
IMF_URLS = ["https://www.imf.org/external/datamapper/api/v2/%s/%s",
            "https://www.imf.org/external/datamapper/api/v1/%s/%s"]

BOC_SERIES = "BD.CDN.10YR.DQ.YLD"
BOC_URL = "https://www.bankofcanada.ca/valet/observations/%s/json" % BOC_SERIES

BUNDESBANK_KEY = "BBSSY/D.REN.EUR.A630.000000WT1010.A"
BUNDESBANK_URL = "https://api.statistiken.bundesbank.de/rest/data/" + BUNDESBANK_KEY

BOE_URL = "https://www.bankofengland.co.uk/boeapps/database/_iadb-fromshowcolumns.asp"
BOE_SERIES = "IUDMNPY"

MOF_URL = "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/jgbcme.csv"

RBA_URL = "https://www.rba.gov.au/statistics/tables/csv/f2-data.csv"
RBA_SERIES = ("FCMYGBAG10D", "FCMYGBAG10")

SNB_CUBE_URL = "https://data.snb.ch/api/cube/%s/data/csv/en"
SNB_YIELD_CUBE = "rendoblid"
SNB_M2_CUBE = "snbmonagg"
SNB_BALANCE_SHEET_CUBE = "snbbipo"
SNB_TEN_YEAR_CODES = ("10J", "10Y", "10", "10 years", "10 Jahre")
SNB_TEN_YEAR_CODE = None   # set to "10J0" or "10J1" once checked against the SNB page (see snb_ten_year_rows)

RBNZ_B2_URL = ("https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily.xlsx")

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def days_before(today, n):
    return today - datetime.timedelta(days=n)


def fetch_cad_yield(today):
    text = http_get_text(BOC_URL, params={"recent": 20})
    points = parse_boc_valet(text, BOC_SERIES)
    return yield_facts("CAD", points, "Bank of Canada Valet API, series %s (10-year benchmark bond yield, daily)"
                       % BOC_SERIES, today)


def fetch_eur_yield(today):
    text = http_get_text(BUNDESBANK_URL, params={"format": "csv", "lang": "en",
                                                 "startPeriod": days_before(today, 30).isoformat()})
    return yield_facts("EUR", parse_sdmx_csv(text),
                       "Deutsche Bundesbank, series BBSSY.D.REN.EUR.A630.000000WT1010.A (daily yield of the "
                       "current 10-year federal bond)", today)


def fetch_gbp_yield(today):
    start = days_before(today, 30)
    params = {"csv.x": "yes", "Datefrom": "%02d/%s/%d" % (start.day, MONTHS[start.month - 1], start.year),
              "Dateto": "now", "SeriesCodes": BOE_SERIES, "CSVF": "TN", "UsingCodes": "Y", "VPD": "Y", "VFD": "N"}
    text = http_get_text(BOE_URL, params=params)
    return yield_facts("GBP", parse_boe_iadb(text, BOE_SERIES),
                       "Bank of England database, series %s (10-year nominal par yield on British government "
                       "securities, daily)" % BOE_SERIES, today)


def fetch_jpy_yield(today):
    return yield_facts("JPY", parse_mof_jgb(http_get_text(MOF_URL)),
                       "Japan Ministry of Finance, JGB interest rate table, 10-year column (daily)", today)


def fetch_aud_yield(today):
    return yield_facts("AUD", parse_rba_table(http_get_text(RBA_URL), RBA_SERIES),
                       "Reserve Bank of Australia, table F2, Australian Government 10 year bond (daily)", today)


RBNZ_HEADERS = {
    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.8",
    "Accept-Language": "en-NZ,en;q=0.9",
    "Referer": "https://www.rbnz.govt.nz/statistics/series/exchange-and-interest-rates/wholesale-interest-rates",
}


def fetch_nzd_yield(today):
    return yield_facts("NZD", parse_rbnz_b2(http_get_bytes(RBNZ_B2_URL, headers=RBNZ_HEADERS)),
                       "Reserve Bank of New Zealand, table B2, 10-year government bond yield (daily)", today)


def snb_cube_text(cube, from_date):
    return http_get_text(SNB_CUBE_URL % cube, params={"fromDate": from_date})


def snb_ten_year_rows(dims, rows):
    """The cube's 10-year rows. The first live run (16 Oct 2026) showed two 10-year codes, 10J0 and 10J1.
    If SNB_TEN_YEAR_CODE is set, that one is used. Otherwise, when more than one code starts with 10J,
    nothing is guessed: the error lists each code's latest value so the right one can be chosen by comparing
    with the SNB's own published 10-year spot rate."""
    if SNB_TEN_YEAR_CODE:
        return snb_series(dims, rows, [SNB_TEN_YEAR_CODE], "the 10-year Swiss Confederation bond yield")
    codes = sorted({r[dim] for dim in dims for r in rows if r.get(dim, "").startswith("10J")})
    if len(codes) <= 1:
        return snb_series(dims, rows, codes or SNB_TEN_YEAR_CODES, "the 10-year Swiss Confederation bond yield")
    latest = {}
    for code in codes:
        mine = sorted((r["Date"], r["Value"]) for r in rows if code in [r.get(dim) for dim in dims])
        latest[code] = mine[-1] if mine else None
    raise FetchError("the cube has several 10-year series and which one is the SNB's 10-year spot rate is not "
                     "settled yet. Latest value of each (date, value): %s" % latest)


def fetch_chf_yield(today):
    dims, rows = parse_snb_cube(snb_cube_text(SNB_YIELD_CUBE, days_before(today, 30).isoformat()))
    picked = snb_ten_year_rows(dims, rows)
    points = [(parse_date(r["Date"], ["%Y-%m-%d"]), r["Value"]) for r in picked]
    return yield_facts("CHF", points,
                       "Swiss National Bank data portal, cube %s (10-year Swiss Confederation bond spot rate, daily)"
                       % SNB_YIELD_CUBE, today)


def snb_monthly_points(cube, code, what, today):
    dims, rows = parse_snb_cube(snb_cube_text(cube, "%04d-%02d" % (today.year - 2, today.month)))
    picked = snb_series(dims, rows, [code], what)
    out = []
    for r in picked:
        try:
            dt = parse_date(r["Date"], ["%Y-%m", "%Y-%m-%d"])
        except FetchError:
            continue
        out.append(((dt.year, dt.month), r["Value"]))
    return out


def snb_m2_points(dims, rows):
    """Swiss M2 monthly levels from the snbmonagg cube. The first live run showed the cube holds the aggregate
    in one dimension (GM1, GM2, GM3 ... ) and a second dimension with two kinds of series (B and VV). The level
    series is the one with the large numbers; the other is expected to be the year-on-year per cent change, and it
    is used as a cross-check: if it disagrees with the change worked out from the levels, nothing is used."""
    agg_dim = next((dim for dim in dims if any(r.get(dim) == "GM2" for r in rows)), None)
    if agg_dim is None:
        raise FetchError("no GM2 (M2) in the cube; it holds %s" % {dim: sorted({r.get(dim) for r in rows})[:30] for dim in dims})
    picked = [r for r in rows if r.get(agg_dim) == "GM2"]
    other = [dim for dim in dims if dim != agg_dim]
    groups = {}
    for r in picked:
        try:
            dt = parse_date(r["Date"], ["%Y-%m", "%Y-%m-%d"])
            groups.setdefault(tuple(r.get(dim) for dim in other), {})[(dt.year, dt.month)] = to_decimal(r["Value"])
        except FetchError:
            continue
    if not groups:
        raise FetchError("no monthly M2 values found")
    ranked = sorted(groups.items(), key=lambda kv: abs(kv[1][max(kv[1])]), reverse=True)
    level_key, level = ranked[0]
    latest_key = max(level)
    if abs(level[latest_key]) < 1000:
        raise FetchError("the largest M2 series is only %s, so it does not look like a level in millions; series seen: %s"
                         % (plain(level[latest_key]), {k: plain(v[max(v)]) for k, v in groups.items()}))
    year_ago = (latest_key[0] - 1, latest_key[1])
    if len(ranked) > 1 and year_ago in level and level[year_ago] != 0:
        computed = (level[latest_key] - level[year_ago]) / level[year_ago] * 100
        checks = []
        for key, series in ranked[1:]:
            if latest_key in series and abs(series[latest_key]) <= 100:
                checks.append((key, series[latest_key]))
        if checks and not any(abs(v - computed) <= Decimal("0.3") for _k, v in checks):
            raise FetchError("the M2 levels give %s%% for %04d-%02d but the cube's other series say %s, so the series "
                             "are not what was assumed and nothing is used" % (
                                 plain(computed.quantize(Decimal("0.1"))), latest_key[0], latest_key[1],
                                 {k: plain(v) for k, v in checks}))
    return sorted(level.items())


def fetch_chf_m2(today):
    dims, rows = parse_snb_cube(snb_cube_text(SNB_M2_CUBE, "%04d-%02d" % (today.year - 2, today.month)))
    pts = snb_m2_points(dims, rows)
    return yoy_facts("CHF", M2_NAME, pts, "Swiss National Bank data portal, cube %s, aggregate M2" % SNB_M2_CUBE,
                     today, "million francs", 75, "Swiss M2")


def fetch_chf_balance_sheet(today):
    pts = snb_monthly_points(SNB_BALANCE_SHEET_CUBE, "T0", "SNB total assets", today)
    return yoy_facts("CHF", BALANCE_SHEET_NAME, pts,
                     "Swiss National Bank data portal, cube %s, total assets (month-end)" % SNB_BALANCE_SHEET_CUBE,
                     today, "billion francs", 75, "SNB total assets", display_divisor=Decimal(1000))


def make_imf_fetcher(currency, indicator):
    series = IMF_SERIES[indicator]
    country = IMF_COUNTRY[currency]

    def run(today):
        last_error = None
        params = {"periods": "%d,%d" % (today.year - 2, today.year - 1)}
        for pattern in IMF_URLS:
            url = pattern % (series, country)
            try:
                values = parse_imf_datamapper(http_get_text(url, params=params), country)
                return imf_facts(currency, indicator, values,
                                 "IMF DataMapper, series %s, country %s (%s)" % (series, country, url), today)
            except Exception as exc:   # try the other address before giving up
                last_error = "%s: %s" % (url, exc)
        raise FetchError(last_error or "no IMF address worked")
    return run


def build_units():
    """Every download the daily fetcher makes, as [{currency, indicator, label, run}]."""
    units = [
        {"currency": "EUR", "indicator": YIELD_NAME, "label": "Bundesbank 10-year Bund yield", "run": fetch_eur_yield},
        {"currency": "GBP", "indicator": YIELD_NAME, "label": "Bank of England IUDMNPY", "run": fetch_gbp_yield},
        {"currency": "JPY", "indicator": YIELD_NAME, "label": "Japan MOF JGB table", "run": fetch_jpy_yield},
        {"currency": "CHF", "indicator": YIELD_NAME, "label": "SNB data portal rendoblid", "run": fetch_chf_yield},
        {"currency": "CAD", "indicator": YIELD_NAME, "label": "Bank of Canada Valet 10-year yield",
         "run": fetch_cad_yield},
        {"currency": "AUD", "indicator": YIELD_NAME, "label": "RBA table F2", "run": fetch_aud_yield},
        {"currency": "NZD", "indicator": YIELD_NAME, "label": "RBNZ table B2", "run": fetch_nzd_yield},
        {"currency": "CHF", "indicator": M2_NAME, "label": "SNB data portal snbmonagg", "run": fetch_chf_m2},
        {"currency": "CHF", "indicator": BALANCE_SHEET_NAME, "label": "SNB data portal snbbipo",
         "run": fetch_chf_balance_sheet},
    ]
    for currency in IMF_COUNTRY:
        for indicator in (DEBT_NAME, FISCAL_NAME):
            units.append({"currency": currency, "indicator": indicator,
                          "label": "IMF DataMapper %s %s" % (IMF_SERIES[indicator], IMF_COUNTRY[currency]),
                          "run": make_imf_fetcher(currency, indicator)})
    return units


def fetch_all(today=None, units=None, log=print, budget_seconds=TOTAL_BUDGET_SECONDS):
    """Run every unit. Returns (facts, failures). A failing unit never stops the others.
    failures = [{currency, indicator, label, reason}]."""
    today = today or datetime.datetime.now(datetime.timezone.utc).date()
    facts, failures = [], []
    started = time.monotonic()
    for unit in (units if units is not None else build_units()):
        try:
            if time.monotonic() - started > budget_seconds:
                raise FetchError("skipped: the downloads ran out of their %d seconds" % budget_seconds)
            got = unit["run"](today)
            if not got:
                raise FetchError("the reader returned nothing")
            facts.extend(got)
            first = got[0]
            log("  OK   %s %s: %s (%s)" % (unit["currency"], unit["indicator"], first["value"], first["as_of"]))
        except Exception as exc:  # noqa: BLE001 - any failure here must fall back to the AI searching
            reason = "%s: %s" % (type(exc).__name__, exc) if not isinstance(exc, FetchError) else str(exc)
            failures.append({"currency": unit["currency"], "indicator": unit["indicator"],
                             "label": unit["label"], "reason": reason})
            log("  FAIL %s %s (%s): %s" % (unit["currency"], unit["indicator"], unit["label"], reason))
            log("::warning::Direct fetch failed for %s %s - the Hub will search for it instead. %s"
                % (unit["currency"], unit["indicator"], reason[:300]))
    return facts, failures


# --------------------------------------------------------------------------------------------------
# The DIRECT AUTO tab: rows to write, and reading them back
# --------------------------------------------------------------------------------------------------

def rows_for_sheet(facts, failures, now):
    """Rows for the DIRECT AUTO tab (below the two header rows). `now` is a UTC datetime."""
    stamp = now.strftime(RETRIEVED_FORMAT)
    rows = []
    for f in facts:
        rows.append([f["currency"], f["indicator"], f["name"], f["value"], f["as_of"], f["source"],
                     STATUS_OK, stamp, ""])
    for fail in failures:
        rows.append([fail["currency"], fail["indicator"], fail["label"], "N/A", "", "", STATUS_FAILED, "",
                     fail["reason"][:300]])
    order = {c: i for i, c in enumerate(["EUR", "GBP", "JPY", "CHF", "CAD", "AUD", "NZD", "USD"])}
    rows.sort(key=lambda r: (order.get(r[0], 99), r[1] != YIELD_NAME, r[1]))
    return rows


def facts_from_rows(rows, now, max_age_days=MAX_FACT_AGE_DAYS):
    """Read DIRECT AUTO rows back. Returns {currency: {fact name: (value, as_of, source)}}.
    Rows that are not OK, or were retrieved more than max_age_days ago, are skipped."""
    out = {}
    for row in rows:
        if len(row) < 8 or str(row[6]).strip().upper() != STATUS_OK:
            continue
        currency, _indicator, name, value, as_of, source = [str(c).strip() for c in row[:6]]
        try:
            retrieved = datetime.datetime.strptime(str(row[7]).strip(), RETRIEVED_FORMAT)
        except ValueError:
            continue
        if not (currency and name and value):
            continue
        if (now.replace(tzinfo=None) - retrieved) > datetime.timedelta(days=max_age_days):
            continue
        out.setdefault(currency, {})[name] = (value, as_of, source)
    return out


if __name__ == "__main__":
    # Dry run for a person at a terminal: downloads everything, prints what it found, writes nothing.
    got, bad = fetch_all()
    for fact in got:
        print("%s | %s | %s | %s" % (fact["currency"], fact["name"], fact["value"], fact["as_of"]))
    print("%d facts, %d failures" % (len(got), len(bad)))
